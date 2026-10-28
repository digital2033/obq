"""Creates templates/question-bank-template.xlsx  (run:  python3 tools/make_template.py)"""
import os
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "templates", "question-bank-template.xlsx")

OT = "Genesis|Exodus|Leviticus|Numbers|Deuteronomy|Joshua|Judges|Ruth|1 Samuel|2 Samuel|1 Kings|2 Kings|1 Chronicles|2 Chronicles|Ezra|Nehemiah|Esther|Job|Psalms|Proverbs|Ecclesiastes|Song of Solomon|Isaiah|Jeremiah|Lamentations|Ezekiel|Daniel|Hosea|Joel|Amos|Obadiah|Jonah|Micah|Nahum|Habakkuk|Zephaniah|Haggai|Zechariah|Malachi".split("|")
NT = "Matthew|Mark|Luke|John|Acts|Romans|1 Corinthians|2 Corinthians|Galatians|Ephesians|Philippians|Colossians|1 Thessalonians|2 Thessalonians|1 Timothy|2 Timothy|Titus|Philemon|Hebrews|James|1 Peter|2 Peter|1 John|2 John|3 John|Jude|Revelation".split("|")
LANGS = ["English", "Hindi", "Malayalam", "Odia", "Tamil", "Telugu", "Punjabi", "Bengali", "Nepali", "Urdu", "Marathi"]
CATS = ["One answer", "Multiple choice", "Short answer with hints", "Clue based", "Multi-part"]
LEVELS = ["Regular", "Trivia", "Difficult"]
HEAD = ["Language", "Book", "Category", "Level", "Question", "Option A", "Option B", "Option C", "Option D",
        "Clue 1", "Clue 2", "Clue 3", "Clue 4", "Clue 5", "Correct answer", "Reference", "More to know", "Points", "Seconds", "Intro"]
WIDTH = [12, 16, 22, 11, 52, 22, 22, 22, 22, 30, 30, 30, 30, 30, 30, 20, 44, 8, 9, 36]
COL = {h: get_column_letter(i + 1) for i, h in enumerate(HEAD)}

ORANGE = "F7931E"; DARK = "1B1410"; CREAM = "FFF6EC"; INPUT = "FFFDF8"
F = "Arial"
thin = Side(style="thin", color="D9C3AE")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ---------------- Instructions ----------------
ws = wb.active
ws.title = "Instructions"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 100
ws["A1"] = "Question bank template"
ws["A1"].font = Font(name=F, size=18, bold=True, color=ORANGE)
ws["A2"] = "Type your questions on the 'Questions' sheet, save the file, then upload it in the Oral Bible Quiz Builder (the 'Quiz Builder' button)."
ws["A2"].font = Font(name=F, size=11, color=DARK)
rows = [
    ("HOW TO USE", None),
    ("1", "Open the 'Examples' sheet to see how each question type is filled in. The Examples sheet is NOT uploaded."),
    ("2", "Go to the 'Questions' sheet. Row 1 is the heading row - do not change the heading names. Type one question per row from row 2."),
    ("3", "Save as .xlsx (or .csv) and upload it in the Oral Bible Quiz Builder. Rows with a problem are skipped and listed on screen with their row number."),
    ("", ""),
    ("QUESTION TYPES (Category)", None),
    ("One answer", "A normal question, no options and no clues. Fill: Question, Correct answer."),
    ("Multiple choice", "Four options, one right answer. Fill: Question, Option A to D, Correct answer (the option text or its letter A to D). The quiz can add a 50:50 button to some of these: the host presses it and two wrong options disappear."),
    ("Short answer with hints", "A one-word or short answer that has hints. Fill: Question, Correct answer, and Clue 1 to Clue 5 as hints shown one at a time. If all clues are blank, hints are made automatically (first letter, number of letters)."),
    ("Clue based", "'Who am I?' style. Fill: Correct answer and Clue 1 to Clue 5, hardest first. Points fall with each clue shown (5 points after clue 1, 4 after clue 2, ...). Question may stay blank. If all clues are blank, clues are made automatically."),
    ("Multi-part", "2 to 4 small questions asked together. In Question write the parts separated by | and in Correct answer the matching answers separated by | in the same order. Optional Intro line. Points = number of parts answered correctly (2 parts: up to 2 points)."),
    ("", ""),
    ("COLUMNS", None),
    ("Language", "Language of the question (pick from the list, or type any other language name such as Kannada). Leave blank if the question can be used in any language."),
    ("Book", "Bible book (pick from the list, or leave blank for a general question). For non-Bible quizzes you may type your own topic, for example 'Church history'."),
    ("Level", "Regular (blank is the same), Trivia (a fun fact) or Difficult. A quiz takes one Trivia and one Difficult question for every 10 questions."),
    ("Points / Seconds", "Optional whole numbers. Blank = 1 point and 30 seconds (45 for Multi-part)."),
    ("Reference / More to know", "Optional. Reference, for example John 3:16. 'More to know' is a short note shown under the answer."),
    ("", ""),
    ("HOW A QUIZ IS BUILT", None),
    ("The mix", "A 10-question quiz is built in this order: 1 regular, 2 multiple choice, 1 trivia, 2 multiple choice with 50:50, 1 short answer with hints, 1 who am I, 1 multi-part, 1 difficult. 20 and 30 questions repeat the mix. If your bank has no question of a kind, the closest kind is used instead, so a bank with only 'One answer' rows still works."),
    ("Scope", "For a Bible quiz, the app picks questions whose Book matches the books chosen on the setup page. Blank Book = general question (used for 'Full Bible' and 'Any topic')."),
    ("Language", "The app uses only the rows that match the Quiz language chosen on the setup page (or rows with a blank Language)."),
    ("Colour key", "Orange cells are headings. Cream cells on the Questions sheet are for you to fill in."),
]
r = 4
for a, b in rows:
    ca, cb = ws.cell(row=r, column=1, value=a), ws.cell(row=r, column=2, value=b)
    if b is None:
        ca.font = Font(name=F, size=12, bold=True, color=DARK)
        for c in (ca, cb):
            c.fill = PatternFill("solid", fgColor=ORANGE)
    else:
        ca.font = Font(name=F, size=11, bold=True, color=DARK)
        cb.font = Font(name=F, size=11, color=DARK)
    ca.alignment = Alignment(vertical="top", wrap_text=True)
    cb.alignment = Alignment(vertical="top", wrap_text=True)
    r += 1

# ---------------- Lists ----------------
wl = wb.create_sheet("Lists")
for i, (title, vals) in enumerate([("Languages", LANGS), ("Categories", CATS), ("Books", ["General"] + OT + NT), ("Levels", LEVELS)]):
    c = wl.cell(row=1, column=i + 1, value=title)
    c.font = Font(name=F, bold=True, color=DARK); c.fill = PatternFill("solid", fgColor=ORANGE)
    for j, v in enumerate(vals):
        wl.cell(row=j + 2, column=i + 1, value=v).font = Font(name=F)
    wl.column_dimensions[get_column_letter(i + 1)].width = 24

# ---------------- Questions ----------------
wq = wb.create_sheet("Questions", 1)
MAXROW = 1000
for ci, h in enumerate(HEAD, 1):
    c = wq.cell(row=1, column=ci, value=h)
    c.font = Font(name=F, bold=True, color=DARK); c.fill = PatternFill("solid", fgColor=ORANGE)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = box
    wq.column_dimensions[get_column_letter(ci)].width = WIDTH[ci - 1]
wq.row_dimensions[1].height = 30
wq.freeze_panes = "A2"
for rr in range(2, 62):
    for ci in range(1, len(HEAD) + 1):
        c = wq.cell(row=rr, column=ci)
        c.fill = PatternFill("solid", fgColor=INPUT); c.border = box
        c.font = Font(name=F, size=11); c.alignment = Alignment(vertical="top", wrap_text=True)
notes = {
    "Category": "Pick: One answer, Multiple choice, Short answer with hints, Clue based or Multi-part.",
    "Level": "Regular (or blank), Trivia or Difficult.",
    "Correct answer": "Multiple choice: exact option text or its letter A-D.  Multi-part: answers separated by |.",
    "Clue 1": "Short answer with hints and Clue based: hints or clues, one per cell. Leave all blank for automatic ones.",
    "Intro": "Multi-part only: a short line shown above the parts.",
}
for ci, h in enumerate(HEAD, 1):
    if h in notes:
        wq.cell(row=1, column=ci).comment = Comment(notes[h], "Template")

def dv(formula, rng, **kw):
    d = DataValidation(type="list", formula1=formula, allow_blank=True, **kw)
    d.error = "Please pick a value from the list."; d.errorTitle = "Not in the list"
    wq.add_data_validation(d); d.add(rng)

dv("=Lists!$B$2:$B$6", "C2:C%d" % MAXROW)
dv("=Lists!$D$2:$D$4", "D2:D%d" % MAXROW)
lg = DataValidation(type="list", formula1="=Lists!$A$2:$A$12", allow_blank=True, showErrorMessage=False)
wq.add_data_validation(lg); lg.add("A2:A%d" % MAXROW)
bk = DataValidation(type="list", formula1="=Lists!$C$2:$C$68", allow_blank=True, showErrorMessage=False)
wq.add_data_validation(bk); bk.add("B2:B%d" % MAXROW)       # custom topics are allowed, so no error message
pv = DataValidation(type="whole", operator="between", formula1="1", formula2="50", allow_blank=True)
pv.error = "Enter a whole number from 1 to 50."; wq.add_data_validation(pv); pv.add("%s2:%s%d" % (COL["Points"], COL["Points"], MAXROW))
sv = DataValidation(type="whole", operator="between", formula1="5", formula2="600", allow_blank=True)
sv.error = "Enter seconds from 5 to 600."; wq.add_data_validation(sv); sv.add("%s2:%s%d" % (COL["Seconds"], COL["Seconds"], MAXROW))

# ---------------- Examples ----------------
we = wb.create_sheet("Examples", 2)
for ci, h in enumerate(HEAD, 1):
    c = we.cell(row=1, column=ci, value=h)
    c.font = Font(name=F, bold=True, color=DARK); c.fill = PatternFill("solid", fgColor=ORANGE)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = box
    we.column_dimensions[get_column_letter(ci)].width = WIDTH[ci - 1]
we.freeze_panes = "A2"
def exrow(lang, book, cat, level, q, opts=(), clues=(), ans="", ref="", info="", pts="", secs="", intro=""):
    o = list(opts) + [""] * (4 - len(opts)); c = list(clues) + [""] * (5 - len(clues))
    return [lang, book, cat, level, q] + o + c + [ans, ref, info, pts, secs, intro]
ex = [
    exrow("English", "John", "One answer", "Regular", "What was the first miracle Jesus performed?", ans="Turning water into wine at Cana", ref="John 2:1-11", info="Six stone water jars were filled and the water became wine."),
    exrow("English", "Genesis", "Multiple choice", "Regular", "Who was the first son of Adam and Eve?", opts=["Cain", "Abel", "Seth", "Enoch"], ans="Cain", ref="Genesis 4:1"),
    exrow("English", "2 Kings", "Short answer with hints", "Regular", "How old was King Josiah when he began to reign?", clues=["Under 20", "A single digit"], ans="Eight", ref="2 Kings 22:1"),
    exrow("English", "Acts", "Clue based", "Regular", "Who am I?", clues=["I was born in the city of Tarsus.", "I studied under the teacher Gamaliel.", "I approved of Stephen's stoning.", "A bright light changed my life on the road to Damascus.", "Jesus sent me to the Gentiles."], ans="Paul", ref="Acts 9; 22", secs=30),
    exrow("English", "Ruth", "Clue based", "Regular", "", ans="Boaz", ref="Ruth 4", info="Clues are made automatically because all five clue cells are blank."),
    exrow("English", "Genesis", "Multi-part", "Regular", "How many days did the rain fall? | Which bird returned with an olive leaf?", ans="Forty days and nights | A dove", ref="Genesis 7:12; 8:11", intro="Noah and the ark: answer both parts"),
    exrow("English", "Numbers", "One answer", "Trivia", "Which animal spoke to Balaam?", ans="His donkey", ref="Numbers 22:28"),
    exrow("English", "Genesis", "One answer", "Difficult", "Who was the father of Abraham?", ans="Terah", ref="Genesis 11:26"),
    exrow("Hindi", "Psalms", "Multiple choice", "Regular", "बाइबल का सबसे लंबा अध्याय कौन सा है?", opts=["भजन संहिता 23", "भजन संहिता 117", "भजन संहिता 119", "यशायाह 53"], ans="C", ref="भजन संहिता 119", info="The answer can be the letter of the right option.", pts=2, secs=40),
]
for ri, row in enumerate(ex, 2):
    for ci, v in enumerate(row, 1):
        c = we.cell(row=ri, column=ci, value=v if v != "" else None)
        c.font = Font(name=F, size=11, italic=True, color="5A4B3E"); c.border = box
        c.alignment = Alignment(vertical="top", wrap_text=True)

for sh in (wq, we):
    sh.sheet_properties.tabColor = ORANGE
wb.active = 1
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("saved", OUT, os.path.getsize(OUT), "bytes")
